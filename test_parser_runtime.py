"""Opt-in real Docling Office conversion: no downloads, OCR models or AI calls.

RUN_PARSER_SMOKE=1 DOCLING_PYTHON=/path/to/parser/python python -m unittest test_parser_runtime
Ordinary CI unit tests skip this; the parser-runtime job installs the heavy runtime.
PDF/OCR model quality is evaluated separately with local rendered-document probes.
"""

import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from document_parser import parse_document


@unittest.skipUnless(
    os.environ.get("RUN_PARSER_SMOKE") == "1", "Optional installed Docling runtime"
)
class InstalledParserTests(unittest.TestCase):
    def test_office_document_preserves_unicode_and_numeric_strings(self):
        # Author a small test fixture, never a production/source document.
        from docx import Document

        with tempfile.TemporaryDirectory() as name:
            path = Path(name) / "fixture.docx"
            original = Document()
            original.add_heading("Dữ liệu giáo dục", level=1)
            original.add_paragraph("Tiếng Việt: phạm vi, đơn vị và nguồn cần được kiểm tra.")
            table = original.add_table(rows=2, cols=2)
            for row, values in zip(
                table.rows, [["Year", "Amount"], ["2024", "999999999999.0001"]], strict=True
            ):
                for cell, text in zip(row.cells, values, strict=True):
                    cell.text = text
            original.save(path)
            with patch.dict(os.environ, {"DOCUMENT_PARSER": "docling", "DOCUMENT_OCR": "off"}):
                parsed = parse_document(
                    path.read_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            self.assertIn("Dữ liệu giáo dục", parsed["text"])
            self.assertIn("999999999999.0001", parsed["text"])
            self.assertEqual(parsed["tables"][0]["cells"][-1]["text"], "999999999999.0001")
            self.assertEqual(parsed["parser"], "docling-local-v1")
            self.assertFalse(parsed["parse_partial"])


if __name__ == "__main__":
    unittest.main()
